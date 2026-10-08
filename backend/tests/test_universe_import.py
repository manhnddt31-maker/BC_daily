from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from backend.app.universe.importer import load_universe_records


def _write_valid_universe(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách cổ phiếu"
    sheet.append(["DANH SÁCH CỔ PHIẾU THEO NHÓM NGÀNH"])
    sheet.append([])
    sheet.append(["STT", "Nhóm ngành", "Mã CP", "Tên công ty", "Sàn"])
    sheet.append([1, "Công nghệ", "FPT", "Công ty Cổ phần FPT", "HOSE"])
    sheet.append([2, "Ngân hàng", "ACB", "Ngân hàng Á Châu", "HOSE"])
    workbook.save(path)


def test_load_universe_records_when_workbook_valid_returns_normalized_rows(
    tmp_path: Path,
) -> None:
    # R1: Tệp hợp lệ phải được chuẩn hóa thành các bản ghi duy nhất.
    workbook_path = tmp_path / "universe.xlsx"
    _write_valid_universe(workbook_path)

    records = load_universe_records(workbook_path)

    assert [record.symbol for record in records] == ["FPT", "ACB"]


def test_load_universe_records_when_exchange_invalid_raises_error(
    tmp_path: Path,
) -> None:
    # R3: Chỉ ba sàn HOSE, HNX và UPCOM được phép xuất hiện trong danh mục.
    workbook_path = tmp_path / "invalid-exchange.xlsx"
    _write_valid_universe(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook.active.cell(row=4, column=5, value="OTC")
    workbook.save(workbook_path)

    with pytest.raises(ValueError):
        load_universe_records(workbook_path)


def test_load_universe_records_when_required_value_missing_raises_error(
    tmp_path: Path,
) -> None:
    # R4: Dòng có mã nhưng thiếu trường bắt buộc phải làm lần import thất bại.
    workbook_path = tmp_path / "missing-required-value.xlsx"
    _write_valid_universe(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook.active.cell(row=4, column=2).value = None
    workbook.save(workbook_path)

    with pytest.raises(ValueError):
        load_universe_records(workbook_path)


def test_load_universe_records_when_no_symbol_rows_raises_error(
    tmp_path: Path,
) -> None:
    # R5: Danh mục không có mã nào không được phép hoàn tất import.
    workbook_path = tmp_path / "empty-universe.xlsx"
    _write_valid_universe(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    sheet.delete_rows(4, 2)
    workbook.save(workbook_path)

    with pytest.raises(ValueError):
        load_universe_records(workbook_path)


def test_load_universe_records_when_symbol_duplicated_raises_error(
    tmp_path: Path,
) -> None:
    # R2: Hai dòng cùng mã phải làm toàn bộ lần import bị từ chối.
    workbook_path = tmp_path / "duplicated.xlsx"
    _write_valid_universe(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook.active.append([3, "Công nghệ", "fpt", "FPT trùng", "HOSE"])
    workbook.save(workbook_path)

    with pytest.raises(ValueError, match="trùng"):
        load_universe_records(workbook_path)
